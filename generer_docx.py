"""
Génère rapport_docusense.docx — conforme Charte IFAG (Licence LMI 2025-2026)
Times New Roman, marges 2.5 cm, interligne 1.5, pagination bas-droite, styles IFAG.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ─── Helpers ──────────────────────────────────────────────────────────────────

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rPr.insert(0, rFonts)

def para_format(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=6, space_after=6, line_spacing=1.5,
                first_indent=None, keep_together=False):
    pf = para.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if keep_together:
        pf.keep_together = True

def add_para(doc, text, font_size=12, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=None,
             space_before=6, space_after=6, first_indent=1.0):
    p = doc.add_paragraph()
    para_format(p, align=align, space_before=space_before,
                space_after=space_after, first_indent=first_indent)
    run = p.add_run(text)
    set_font(run, size=font_size, bold=bold, italic=italic, color=color)
    return p

def add_heading(doc, text, level=1):
    """Niveaux: 1=Chapitre 14pt gras, 2=Section 13pt gras, 3=Sous-section 12pt gras-italique"""
    p = doc.add_paragraph()
    if level == 1:
        para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=12, first_indent=0)
        run = p.add_run(text)
        set_font(run, size=14, bold=True)
        p.style = doc.styles['Heading 1']
    elif level == 2:
        para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6, first_indent=0)
        run = p.add_run(text)
        set_font(run, size=13, bold=True)
        p.style = doc.styles['Heading 2']
    elif level == 3:
        para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6, first_indent=0)
        run = p.add_run(text)
        set_font(run, size=12, bold=True, italic=True)
        p.style = doc.styles['Heading 3']
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=2, space_after=2, first_indent=0)
    if bold_prefix:
        r1 = p.add_run(bold_prefix + " ")
        set_font(r1, bold=True)
        r2 = p.add_run(text)
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def page_break(doc):
    doc.add_page_break()

def horizontal_rule(doc):
    p = doc.add_paragraph()
    para_format(p, space_before=2, space_after=2, first_indent=0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '185FA5')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_page_number_footer(section):
    """Ajoute numéro de page en bas à droite."""
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.clear()
    run = fp.add_run()
    set_font(run, size=10)
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def configure_section(section, margin=2.5):
    section.top_margin = Cm(margin)
    section.bottom_margin = Cm(margin)
    section.left_margin = Cm(margin)
    section.right_margin = Cm(margin)

# ─── Document ─────────────────────────────────────────────────────────────────

doc = Document()

# Styles globaux
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
for h_name in ['Heading 1', 'Heading 2', 'Heading 3']:
    hs = doc.styles[h_name]
    hs.font.name = 'Times New Roman'
    hs.font.color.rgb = RGBColor(0, 0, 0)

section = doc.sections[0]
configure_section(section)

# ═══════════════════════════════════════════════════════════════
# PAGE DE GARDE (modèle Annexe 2 fascicule IFAG)
# ═══════════════════════════════════════════════════════════════
section.different_first_page_header_footer = True  # pas de pied sur couverture

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4, first_indent=0)
run = p.add_run("Institut de Formation d'Assurances et de Gestion (IFAG)")
set_font(run, size=14, bold=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("Établissement Privé de Formation Supérieure")
set_font(run, size=11, italic=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("Agréé par le Ministère de l'Enseignement Supérieur et de la Recherche Scientifique")
set_font(run, size=10, italic=True)

horizontal_rule(doc)

for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, first_indent=0)
run = p.add_run("MÉMOIRE DE FIN D'ÉTUDES")
set_font(run, size=13, bold=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("Pour l'obtention du diplôme de")
set_font(run, size=12)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("LICENCE")
set_font(run, size=16, bold=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("Filière : Informatique   |   Spécialité : Métiers de l'Informatique")
set_font(run, size=12)

for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, first_indent=0)
run = p.add_run("DocuSense")
set_font(run, size=22, bold=True, color=(24, 95, 165))

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, first_indent=0)
run = p.add_run("Plateforme Intelligente de Génération de Documentation Technique Gamifiée")
set_font(run, size=14, bold=True, italic=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, first_indent=0)
run = p.add_run("Propulsée par l'Intelligence Artificielle Générative (Google Gemini 2.0 Flash)")
set_font(run, size=12, italic=True)

for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4, first_indent=0)
r1 = p.add_run("Réalisé par : ")
set_font(r1, size=12, bold=True)
r2 = p.add_run("Boumehdi Rym")
set_font(r2, size=12)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4, first_indent=0)
r1 = p.add_run("Encadré par : ")
set_font(r1, size=12, bold=True)
r2 = p.add_run("[Prénom et Nom de l'encadrant]")
set_font(r2, size=12)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, first_indent=0)
r1 = p.add_run("Année universitaire : ")
set_font(r1, size=12, bold=True)
r2 = p.add_run("2025 – 2026")
set_font(r2, size=12)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# DÉDICACE
# ═══════════════════════════════════════════════════════════════
add_para(doc, "Dédicace", font_size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=60, space_after=20, first_indent=0)

for ligne in [
    "À mes parents, pour leur soutien indéfectible et leurs sacrifices.",
    "À mes enseignants, pour la transmission du savoir.",
    "À tous ceux qui croient que la technologie peut rendre la connaissance accessible à tous.",
]:
    p = doc.add_paragraph()
    para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=6, space_after=6, first_indent=0)
    run = p.add_run(ligne)
    set_font(run, size=12, italic=True)

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=20, space_after=0, first_indent=0)
run = p.add_run("Boumehdi Rym")
set_font(run, size=12, italic=True)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# REMERCIEMENTS
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Remerciements", level=1)
add_para(doc, "Avant tout, je remercie Allah de m'avoir accordé la santé, la volonté et la persévérance nécessaires pour mener à bien ce projet de fin d'études.")
add_para(doc, "Je tiens à exprimer ma profonde gratitude à mon encadrant, dont les conseils avisés, la disponibilité et l'accompagnement rigoureux ont été déterminants dans la réalisation de ce travail. Ses orientations m'ont permis de structurer ma réflexion et d'approfondir mes connaissances tout au long de ce projet.")
add_para(doc, "Mes remerciements s'adressent également à l'ensemble du corps enseignant du département Informatique de l'IFAG pour la qualité de la formation dispensée au cours de ces trois années de licence.")
add_para(doc, "Je remercie les membres du jury d'avoir accepté d'évaluer ce travail.")
add_para(doc, "Enfin, je remercie ma famille et mes amis pour leur soutien moral, leur patience et leurs encouragements tout au long de mon parcours académique.")

p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=12, space_after=0, first_indent=0)
run = p.add_run("Boumehdi Rym")
set_font(run, size=12, italic=True)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# RÉSUMÉ
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Résumé", level=1)

add_para(doc, "Résumé – Français", font_size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=4, first_indent=0)
add_para(doc, "La documentation technique constitue un défi majeur pour les développeurs logiciels : elle est chronophage, souvent négligée et difficile à maintenir. Ce projet de fin d'études présente DocuSense, une plateforme web intelligente qui automatise la création de manuels utilisateurs grâce à l'intelligence artificielle générative.")
add_para(doc, "DocuSense exploite l'API Google Gemini 2.0 Flash pour analyser des captures d'écran et des vidéos d'interfaces applicatives, puis générer automatiquement des étapes de documentation structurées. La plateforme intègre un système de gamification composé de quiz interactifs, de badges de progression (Bronze, Silver, Gold) et d'un classement compétitif, afin d'améliorer l'engagement et la rétention des connaissances par les utilisateurs finaux.")
add_para(doc, "Techniquement, l'application repose sur une architecture avec React.js pour le frontend, Node.js/Express.js pour le backend, MySQL via Sequelize ORM pour la persistance des données, et une authentification sécurisée par JSON Web Tokens (JWT).")
add_para(doc, "Mots-clés : Documentation automatique, Intelligence Artificielle Générative, Gemini AI, Gamification, JWT.",
         bold=False, italic=True, space_before=4, space_after=10)

horizontal_rule(doc)

add_para(doc, "Abstract – English", font_size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, first_indent=0)
add_para(doc, "Technical documentation represents a major challenge for software developers: it is time-consuming, often neglected, and difficult to maintain. This final year project presents DocuSense, an intelligent web platform that automates the creation of user manuals through generative artificial intelligence.")
add_para(doc, "DocuSense leverages the Google Gemini 2.0 Flash API to analyze screenshots and videos of application interfaces, then automatically generates structured documentation steps. The platform integrates a gamification system consisting of interactive quizzes, progress badges (Bronze, Silver, Gold), and a competitive leaderboard, in order to improve engagement and knowledge retention by end users.")
add_para(doc, "Technically, the application relies on React.js for the frontend, Node.js/Express.js for the backend, MySQL via Sequelize ORM for data persistence, and secure authentication via JSON Web Tokens (JWT).")
add_para(doc, "Keywords: Automatic Documentation, Generative Artificial Intelligence, Gemini AI, Gamification, JWT.",
         bold=False, italic=True, space_before=4, space_after=10)

horizontal_rule(doc)

add_para(doc, "ملخص – العربية", font_size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=10, space_after=4, first_indent=0)
add_para(doc,
    "تُمثِّل الوثائق التقنية تحديًا كبيرًا للمطورين البرمجيين، إذ تستغرق وقتًا طويلاً، وكثيرًا ما يُهمَل إعدادها ويصعب الحفاظ عليها. يُقدِّم هذا المشروع منصة DocuSense، وهي منصة ويب ذكية تعمل على أتمتة إنشاء أدلة المستخدمين بالاستعانة بالذكاء الاصطناعي التوليدي.",
    align=WD_ALIGN_PARAGRAPH.RIGHT, first_indent=0)
add_para(doc,
    "تستعين DocuSense بواجهة برمجة تطبيقات Google Gemini 2.0 Flash لتحليل لقطات الشاشة ومقاطع الفيديو لواجهات التطبيقات، وتوليد خطوات التوثيق المنظمة تلقائيًا. تتضمن المنصة نظام تلعيب مؤلفًا من اختبارات تفاعلية وشارات تقدم وقائمة متصدرين تنافسية.",
    align=WD_ALIGN_PARAGRAPH.RIGHT, first_indent=0)
add_para(doc,
    "تقنيًا، يعتمد التطبيق على React.js للواجهة الأمامية، وNode.js/Express.js للخلفية، وMySQL عبر Sequelize ORM، والمصادقة الآمنة بواسطة JWT.",
    align=WD_ALIGN_PARAGRAPH.RIGHT, first_indent=0)
add_para(doc, "الكلمات المفتاحية: توثيق تلقائي، ذكاء اصطناعي توليدي، تلعيب، Gemini AI، JWT.",
         italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=4, space_after=10, first_indent=0)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# TABLE DES MATIÈRES (placeholder — Google Docs génère automatiquement)
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Table des matières", level=1)
add_para(doc,
    "[Dans Google Docs : Insertion → Table des matières → Choisir le style]",
    italic=True, color=(150, 150, 150), first_indent=0)
page_break(doc)

# ═══════════════════════════════════════════════════════════════
# LISTE DES ABRÉVIATIONS
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Liste des abréviations", level=1)

abbreviations = [
    ("AI / IA", "Artificial Intelligence / Intelligence Artificielle"),
    ("API", "Application Programming Interface"),
    ("CRUD", "Create, Read, Update, Delete"),
    ("HTTP", "HyperText Transfer Protocol"),
    ("JSON", "JavaScript Object Notation"),
    ("JWT", "JSON Web Token"),
    ("LLM", "Large Language Model"),
    ("ORM", "Object-Relational Mapping"),
    ("PFE", "Projet de Fin d'Études"),
    ("REST", "Representational State Transfer"),
    ("SGBDR", "Système de Gestion de Base de Données Relationnelle"),
    ("SPA", "Single Page Application"),
    ("SQL", "Structured Query Language"),
    ("UX", "User Experience"),
]
for abbr, meaning in abbreviations:
    p = doc.add_paragraph()
    para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=2, first_indent=0)
    r1 = p.add_run(f"{abbr} : ")
    set_font(r1, bold=True)
    r2 = p.add_run(meaning)
    set_font(r2)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# INTRODUCTION GÉNÉRALE
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Introduction générale", level=1)

add_heading(doc, "Contexte", level=2)
add_para(doc, "Dans l'écosystème numérique contemporain, la documentation technique occupe une place centrale dans la réussite d'un produit logiciel. Elle constitue le pont entre les développeurs qui créent un logiciel et les utilisateurs finaux qui l'exploitent au quotidien. Pourtant, la réalité du terrain montre que cette documentation est trop souvent sacrifiée sur l'autel des délais de livraison : complexe à maintenir, chronophage à rédiger, elle est reléguée au second plan dans la plupart des cycles de développement.")
add_para(doc, "Une étude menée par la société Stack Overflow auprès de développeurs révèle que la documentation insuffisante ou inexistante est l'une des principales sources de frustration dans les projets logiciels. Face à ce constat, l'intelligence artificielle générative ouvre de nouvelles perspectives : des modèles de langage multimodaux, capables de comprendre à la fois du texte et des images, peuvent désormais analyser automatiquement une interface graphique et en extraire une documentation structurée.")

add_heading(doc, "Problématique", level=2)
add_para(doc, "Comment permettre à un développeur de générer en quelques minutes une documentation utilisateur complète, interactive et engageante, à partir d'une simple capture d'écran ou d'une vidéo de son application ?")

add_heading(doc, "Objectifs", level=2)
add_para(doc, "Ce projet de fin d'études vise à concevoir et développer DocuSense, une plateforme web qui :", first_indent=0)
add_bullet(doc, "Automatise la génération de documentation technique via l'IA générative multimodale ;")
add_bullet(doc, "Intègre un système de gamification (quiz, badges, classement) pour améliorer l'engagement des utilisateurs ;")
add_bullet(doc, "Offre une interface moderne, intuitive et responsive ;")
add_bullet(doc, "Garantit la sécurité des données via une authentification par JWT.")

add_heading(doc, "Structure du rapport", level=2)
add_para(doc, "Ce rapport est organisé en trois chapitres :", first_indent=0)
add_bullet(doc, "État de l'art et étude de l'existant :", "Chapitre 1 –")
add_bullet(doc, "Analyse et conception :", "Chapitre 2 –")
add_bullet(doc, "Réalisation et tests :", "Chapitre 3 –")

add_heading(doc, "Déclaration sur l'usage de l'Intelligence Artificielle", level=2)
add_para(doc, "Conformément à la Charte de rédaction de l'IFAG (§ 2.18), nous déclarons explicitement que des outils d'intelligence artificielle générative ont été utilisés dans le cadre de ce projet :")
add_bullet(doc, "constitue le cœur fonctionnel du moteur d'analyse de DocuSense ; son usage est l'objet même du projet.", "Google Gemini 2.0 Flash")
add_bullet(doc, "ont été ponctuellement employés comme aide au débogage et à la rédaction de certains passages.", "Des assistants IA (Claude, ChatGPT)")
add_para(doc, "L'ensemble du code a été intégralement compris, testé et validé par l'étudiante. Tout contenu généré par IA a été relu, reformulé et contextualisé.")

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# CHAPITRE 1 — ÉTAT DE L'ART
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Chapitre 1 – État de l'art et étude de l'existant", level=1)

add_heading(doc, "Introduction", level=2)
add_para(doc, "La conception d'une plateforme intelligente de génération de documentation nécessite une maîtrise solide du domaine dans lequel elle s'inscrit. Avant d'aborder l'implémentation technique de DocuSense, il est indispensable d'établir un socle conceptuel clair et de situer notre contribution par rapport aux travaux et outils existants. Ce chapitre constitue ce fondement théorique et analytique.")
add_para(doc, "Nous débutons par une définition de la documentation technique logicielle, en examinant ses formes, ses enjeux et les limites des approches manuelles traditionnelles. Nous présentons ensuite les avancées récentes en matière d'intelligence artificielle générative et de traitement multimodal, qui constituent le moteur de notre solution. Une étude comparative des outils existants dans ce domaine nous permettra de dégager le positionnement différenciant de DocuSense. Enfin, nous exposons et justifions l'ensemble des technologies retenues pour le développement de la plateforme.")

add_heading(doc, "La documentation technique logicielle", level=2)

add_heading(doc, "Définition et importance", level=3)
add_para(doc, "La documentation technique logicielle désigne l'ensemble des artefacts écrits ou multimédias qui décrivent l'utilisation, la configuration, l'architecture et les fonctionnalités d'un système informatique. Elle se décline en deux catégories principales : la documentation destinée aux développeurs (API docs, code docs) et la documentation destinée aux utilisateurs finaux (manuels utilisateurs, guides de démarrage).")
add_para(doc, "La documentation utilisateur remplit plusieurs fonctions essentielles :")
add_bullet(doc, "Une documentation claire diminue significativement le volume de tickets d'assistance.", "Réduction du support technique :")
add_bullet(doc, "Les utilisateurs autonomes adoptent plus rapidement un produit.", "Accélération de l'adoption :")
add_bullet(doc, "Une bonne documentation réduit la courbe d'apprentissage.", "Amélioration de l'expérience utilisateur :")
add_bullet(doc, "Elle facilite la maintenance et l'évolution du logiciel.", "Pérennité du produit :")

add_heading(doc, "Les défis de la documentation manuelle", level=3)
add_para(doc, "La rédaction manuelle de documentation présente plusieurs inconvénients majeurs qui freinent son adoption systématique dans les équipes de développement. En termes de coût en temps, la rédaction, la mise à jour et la validation d'un manuel utilisateur complet peuvent représenter plusieurs jours de travail. Concernant la cohérence, maintenir une documentation synchronisée avec les évolutions du logiciel est un défi permanent. Du point de vue de la qualité, la documentation rédigée à la hâte est souvent incomplète ou obsolète. Enfin, quant à l'engagement, les manuels statiques génèrent peu d'interaction et sont rarement lus en intégralité.")

add_heading(doc, "L'Intelligence Artificielle Générative", level=2)

add_heading(doc, "Définition et évolution", level=3)
add_para(doc, "L'Intelligence Artificielle Générative (IA générative) désigne une catégorie de systèmes d'IA capables de produire de nouveaux contenus — texte, images, code, audio, vidéo — à partir d'un apprentissage sur de larges corpus de données. Contrairement aux systèmes d'IA discriminatifs qui classifient ou prédisent, les modèles génératifs apprennent la distribution sous-jacente des données d'entraînement et peuvent en générer de nouvelles instances.")
add_para(doc, "L'avènement des Transformers (Vaswani et al., 2017) et des Large Language Models (LLM) comme GPT, PaLM, et Gemini a révolutionné le domaine. Ces modèles, entraînés sur des milliards de paramètres et de tokens, démontrent des capacités remarquables en compréhension et génération de langage naturel.")

add_heading(doc, "Les modèles multimodaux", level=3)
add_para(doc, "Les modèles multimodaux constituent une évolution majeure des LLM : ils sont capables de traiter simultanément plusieurs modalités d'entrée (texte, images, vidéos, audio). Google Gemini 2.0 Flash, utilisé dans DocuSense, est un modèle multimodal de dernière génération capable d'analyser une capture d'écran et d'en décrire les éléments d'interface de manière structurée.")

add_heading(doc, "Étude de l'existant et analyse comparative", level=2)
add_para(doc, "Plusieurs solutions existent pour automatiser ou faciliter la création de documentation. Parmi les principales, on distingue les outils traditionnels (Confluence, Notion, GitBook), les outils IA spécialisés (Mintlify, Swimm, Scribe), et les LLM généraux (ChatGPT, Gemini). Cependant, aucun de ces outils ne combine nativement l'analyse multimodale d'interfaces avec un système de gamification intégré visant l'engagement des utilisateurs finaux.")
add_para(doc, "DocuSense se positionne comme une solution originale qui comble ce gap en proposant une plateforme unifiée : génération IA depuis screenshots/vidéos + gamification (quiz, badges, classement) + gestion de manuels.")

add_heading(doc, "Technologies retenues", level=2)
add_para(doc, "L'architecture de DocuSense repose sur un ensemble de technologies modernes et éprouvées, sélectionnées pour leur pertinence technique et leur complémentarité.")

add_heading(doc, "Frontend – React.js", level=3)
add_para(doc, "React.js (v19) est une bibliothèque JavaScript développée par Meta pour la construction d'interfaces utilisateur. Ses avantages pour DocuSense sont : le DOM Virtuel pour des performances optimales, l'architecture par composants réutilisables, l'écosystème riche (React Router, Axios), et l'état de l'art en développement frontend.")

add_heading(doc, "Backend – Node.js / Express.js", level=3)
add_para(doc, "Node.js est un environnement d'exécution JavaScript côté serveur, basé sur le moteur V8 de Chrome. Express.js (v5) est un framework minimaliste qui facilite la création d'API RESTful. Cette combinaison offre : un langage unifié (JavaScript) côté client et serveur, des performances élevées grâce au modèle non-bloquant, et une architecture modulaire.")

add_heading(doc, "Base de données – MySQL / Sequelize", level=3)
add_para(doc, "MySQL est un Système de Gestion de Base de Données Relationnelle (SGBDR) open-source, largement utilisé dans les applications web. Sequelize est un ORM (Object-Relational Mapping) qui permet de manipuler la base de données via des objets JavaScript, offrant une abstraction de la couche SQL et une gestion simplifiée des migrations.")

add_heading(doc, "IA – Google Gemini 2.0 Flash", level=3)
add_para(doc, "Gemini 2.0 Flash est le modèle d'IA générative multimodale de Google, optimisé pour la rapidité et l'efficacité. Il est capable d'analyser du texte, des images et des vidéos pour générer des réponses structurées. Son intégration via le SDK @google/generative-ai permet d'envoyer des captures d'écran et d'obtenir en retour des étapes de documentation structurées au format JSON.")

add_heading(doc, "Conclusion", level=2)
add_para(doc, "Ce premier chapitre nous a permis de poser le cadre conceptuel et technologique du projet. L'analyse de l'existant révèle un manque de solutions combinant génération IA multimodale et gamification pour la documentation utilisateur. Les technologies retenues — React.js, Node.js/Express.js, MySQL, et Google Gemini — forment un ensemble cohérent et adapté aux objectifs de DocuSense. Le chapitre suivant présentera la phase d'analyse et de conception de la plateforme.")

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# CHAPITRE 2 — ANALYSE ET CONCEPTION
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Chapitre 2 – Analyse et conception", level=1)

add_heading(doc, "Introduction", level=2)
add_para(doc, "Après avoir établi le contexte théorique et technologique dans le chapitre précédent, ce chapitre est consacré à la phase d'analyse et de conception de DocuSense. Nous y définissons les besoins fonctionnels et non fonctionnels de la plateforme, modélisons les cas d'utilisation via des diagrammes UML, puis présentons l'architecture logicielle et le modèle de données retenu.")

add_heading(doc, "Spécification des besoins", level=2)

add_heading(doc, "Besoins fonctionnels", level=3)
add_para(doc, "Les besoins fonctionnels définissent ce que le système doit faire. Pour DocuSense, ils s'articulent autour de cinq axes :")
add_bullet(doc, "Inscription, connexion et déconnexion sécurisées via JWT.", "Gestion des utilisateurs :")
add_bullet(doc, "Création, lecture, mise à jour et suppression de manuels et d'étapes.", "Gestion des manuels (CRUD) :")
add_bullet(doc, "Génération d'étapes depuis une description texte, une image (screenshot) ou une vidéo via l'API Gemini.", "Module IA :")
add_bullet(doc, "Quiz interactifs par manuel, attribution de badges (Bronze/Silver/Gold), classement Top 10.", "Gamification :")
add_bullet(doc, "Suivi de la progression par utilisateur et par manuel.", "Suivi de progression :")

add_heading(doc, "Besoins non fonctionnels", level=3)
add_para(doc, "Les besoins non fonctionnels définissent les contraintes de qualité du système :")
add_bullet(doc, "Temps de réponse de l'API < 3 secondes pour les requêtes standards.", "Performance :")
add_bullet(doc, "Mots de passe hachés (bcrypt), authentification JWT, protection CORS.", "Sécurité :")
add_bullet(doc, "Interface responsive compatible desktop et mobile.", "Accessibilité :")
add_bullet(doc, "Architecture modulaire permettant l'ajout de nouvelles fonctionnalités.", "Extensibilité :")
add_bullet(doc, "Documentation du code et commentaires pour faciliter la maintenance.", "Maintenabilité :")

add_heading(doc, "Modélisation UML", level=2)

add_heading(doc, "Diagramme des cas d'utilisation", level=3)
add_para(doc,
    "[Figure 2.1 – Diagramme des cas d'utilisation de DocuSense]\n"
    "Deux acteurs principaux : le Développeur (gestion des manuels, module IA) "
    "et l'Utilisateur Final (consultation, quiz, gamification).",
    italic=True, color=(100, 100, 100), first_indent=0)

add_heading(doc, "Diagramme de classes", level=3)
add_para(doc, "Le modèle de données de DocuSense comprend cinq entités principales : User, Manual, Step, Quiz et Progress. La relation centrale est User → Manual (1,N) → Step (1,N). Les entités Quiz et Progress sont associées respectivement aux manuels et aux utilisateurs pour la gestion de la gamification.")

add_heading(doc, "Diagramme de séquence", level=3)
add_para(doc, "Le diagramme de séquence pour la génération IA illustre l'interaction entre l'utilisateur, le frontend React, le backend Express et l'API Gemini : l'utilisateur soumet une image → le frontend envoie une requête multipart → le backend convertit l'image en Base64 → l'API Gemini analyse et retourne des étapes JSON → le backend les renvoie au frontend qui les affiche.")

add_heading(doc, "Architecture logicielle", level=2)

add_heading(doc, "Architecture trois-tiers", level=3)
add_para(doc, "DocuSense repose sur une architecture client-serveur à trois niveaux (three-tier) : le Frontend (React.js) communique via HTTP/REST avec le Backend (Node.js/Express.js), qui interagit avec la Base de données (MySQL) via SQL/ORM et avec l'API Gemini via HTTPS.")

add_heading(doc, "Architecture RESTful", level=3)
add_para(doc, "Les routes API de DocuSense suivent les conventions REST : /api/auth (inscription, connexion), /api/manuals (CRUD manuels et étapes), /api/ai (suggest, analyze-image, analyze-video), /api/quiz (quiz et gamification). Chaque route protégée est sécurisée par un middleware JWT.")

add_heading(doc, "Modèle de données", level=2)
add_para(doc, "La base de données MySQL contient cinq tables principales : users (id, name, email, password, createdAt), manuals (id, userId, title, description, isPublished), steps (id, manualId, title, description, order), quizzes (id, manualId, question, options JSON, answer), progress (id, userId, manualId, score, badge, completedAt).")

add_heading(doc, "Conclusion", level=2)
add_para(doc, "Ce chapitre a présenté l'analyse complète des besoins de DocuSense et sa conception architecturale. Les diagrammes UML ont permis de formaliser les interactions entre les acteurs et le système, tandis que l'architecture trois-tiers et le modèle de données fournissent une base solide pour l'implémentation. Le chapitre suivant détaille la réalisation technique et les interfaces de la plateforme.")

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# CHAPITRE 3 — RÉALISATION ET TESTS
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Chapitre 3 – Réalisation et tests", level=1)

add_heading(doc, "Introduction", level=2)
add_para(doc, "Ce chapitre présente la phase de réalisation de DocuSense : l'environnement de développement utilisé, les aspects techniques clés de l'implémentation, les interfaces graphiques de la plateforme, et les résultats des tests effectués. Il constitue la validation concrète des choix de conception établis au chapitre précédent.")

add_heading(doc, "Environnement de développement", level=2)

add_heading(doc, "Outils et versions", level=3)
add_para(doc, "Le développement de DocuSense a été réalisé avec les technologies et versions suivantes :", first_indent=0)

technologies = [
    ("Node.js 20.x", "Environnement d'exécution serveur"),
    ("React.js 19.2.5", "Framework frontend"),
    ("Express.js 5.2.1", "Framework backend"),
    ("Sequelize 6.37.8", "ORM pour MySQL"),
    ("MySQL 8.x", "Base de données relationnelle"),
    ("@google/generative-ai 0.24.1", "SDK Gemini"),
    ("Multer 2.1.1", "Upload de fichiers"),
    ("Sharp 0.34.5", "Traitement d'images"),
    ("bcryptjs 3.0.3", "Hachage de mots de passe"),
    ("jsonwebtoken 9.0.3", "Authentification JWT"),
    ("VS Code", "Éditeur de code"),
    ("Postman", "Test des API REST"),
]
for tech, role in technologies:
    p = doc.add_paragraph()
    para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=1, space_after=1, first_indent=0)
    r1 = p.add_run(f"{tech} : ")
    set_font(r1, bold=True, size=11)
    r2 = p.add_run(role)
    set_font(r2, size=11)

add_heading(doc, "Implémentation des fonctionnalités clés", level=2)

add_heading(doc, "Authentification JWT", level=3)
add_para(doc, "Le système d'authentification de DocuSense repose sur JSON Web Tokens (JWT). Lors de la connexion, le backend génère un token signé contenant l'identifiant de l'utilisateur et le renvoie au client. Ce token est stocké dans le localStorage du navigateur et envoyé dans l'en-tête Authorization de chaque requête protégée. Le middleware authMiddleware.js vérifie la validité et l'expiration du token avant chaque opération sécurisée.")

add_heading(doc, "Module IA – Génération par description", level=3)
add_para(doc, "La route POST /api/ai/suggest reçoit une description textuelle de l'application cible. Elle construit un prompt structuré demandant à Gemini 2.0 Flash de générer exactement 5 étapes de documentation au format JSON (titre + description par étape). La réponse JSON est parsée et retournée au frontend pour affichage et sauvegarde.")

add_heading(doc, "Module IA – Analyse d'image", level=3)
add_para(doc, "La route POST /api/ai/analyze-image accepte un fichier image (PNG/JPG) via Multer. L'image est traitée par Sharp pour optimisation, puis convertie en Base64. Ce contenu encodé est envoyé à l'API Gemini avec le type MIME approprié (image/png ou image/jpeg). Gemini analyse l'interface graphique et retourne des étapes structurées décrivant les fonctionnalités visibles.")

add_heading(doc, "Système de gamification", level=3)
add_para(doc, "La gamification de DocuSense comprend trois composantes : les quiz (questions à choix multiples générées pour chaque manuel), les badges (Bronze ≥ 50 pts, Silver ≥ 100 pts, Gold ≥ 200 pts attribués selon le score obtenu), et le classement (Top 10 des utilisateurs par score total, mis à jour dynamiquement).")

add_heading(doc, "Présentation des interfaces", level=2)

add_heading(doc, "Page d'accueil (Landing Page)", level=3)
add_para(doc, "La page d'accueil présente la plateforme DocuSense avec une navbar, une section Hero illustrant la génération automatique de documentation, des statistiques de la plateforme, une présentation des fonctionnalités clés, et un appel à l'action. Elle permet la navigation vers les pages d'inscription et de connexion.")

add_heading(doc, "Tableau de bord", level=3)
add_para(doc, "Le tableau de bord est l'interface principale de l'utilisateur authentifié. Il affiche les métriques clés (total manuels, manuels publiés, score total), la liste des manuels récents avec leurs statuts, et les outils de génération IA intégrés directement sur chaque carte manuelle. La sidebar permet la navigation entre les différentes sections.")

add_heading(doc, "Module Quiz et Gamification", level=3)
add_para(doc, "La page Quiz présente les questions associées aux manuels de l'utilisateur. Après complétion, le score est calculé, le badge correspondant est attribué, et la progression est enregistrée. Un classement compétitif Top 10 est affiché pour stimuler l'engagement.")

add_heading(doc, "Tests et validation", level=2)

add_heading(doc, "Tests fonctionnels", level=3)
add_para(doc, "Les tests fonctionnels ont été réalisés avec Postman pour les API REST et manuellement pour le frontend. Les scénarios testés couvrent : l'inscription et la connexion d'un utilisateur, la création et la gestion de manuels, la génération IA par description et par image, la complétion d'un quiz et l'attribution d'un badge, et la consultation du classement.")

add_heading(doc, "Tests de sécurité", level=3)
add_para(doc, "Les tests de sécurité ont vérifié : la protection des routes par JWT (accès refusé sans token valide), le hachage des mots de passe avec bcrypt (vérification en base de données), la protection CORS (seul le frontend autorisé peut appeler l'API), et la non-exposition des variables d'environnement (clés API, secrets JWT).")

add_heading(doc, "Conclusion", level=2)
add_para(doc, "Ce chapitre a présenté la réalisation concrète de DocuSense, depuis l'environnement de développement jusqu'aux interfaces graphiques finales. Les tests fonctionnels et de sécurité ont validé le bon fonctionnement de l'ensemble des fonctionnalités implémentées. DocuSense répond ainsi aux objectifs fixés en introduction : automatisation de la documentation par IA, gamification de l'apprentissage, et interface moderne sécurisée.")

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# CONCLUSION GÉNÉRALE
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Conclusion générale et perspectives", level=1)

add_heading(doc, "Rappel de la problématique et des objectifs", level=2)
add_para(doc, "Ce projet de fin d'études avait pour objectif de répondre à un défi récurrent dans le développement logiciel : la génération de documentation utilisateur de qualité, rapidement et à moindre effort. La problématique centrale était : comment permettre à un développeur de produire une documentation interactive et engageante en quelques minutes, à partir d'une simple capture d'écran ou d'une vidéo ?")

add_heading(doc, "Bilan des réalisations", level=2)
add_para(doc, "DocuSense répond à cette problématique par une plateforme web complète intégrant :")
add_bullet(doc, "Analyse de captures d'écran et de vidéos d'interfaces applicatives via Google Gemini 2.0 Flash pour générer automatiquement des étapes de documentation structurées.", "Un moteur IA multimodal :")
add_bullet(doc, "Quiz interactifs, badges de progression (Bronze, Silver, Gold) et classement compétitif pour transformer la consultation de documentation en expérience engageante.", "Un système de gamification :")
add_bullet(doc, "Une architecture React.js / Node.js / MySQL sécurisée par JWT, offrant une interface moderne et responsive.", "Une plateforme technique robuste :")

add_heading(doc, "Limites et contraintes", level=2)
add_para(doc, "Plusieurs limites ont été identifiées au cours du développement : la dépendance à l'API Google Gemini (quota limité et connectivité requise), l'absence de fonctionnalité d'export de la documentation (PDF, HTML), la non-internationalisation de l'interface (actuellement en français uniquement), et l'absence de tests automatisés (tests unitaires et d'intégration).")

add_heading(doc, "Perspectives", level=2)
add_para(doc, "Les pistes d'amélioration et d'extension envisagées pour DocuSense sont :")
add_bullet(doc, "Export des manuels en PDF, HTML et Markdown pour une diffusion multi-canal.")
add_bullet(doc, "Internationalisation de l'interface (arabe, anglais, français).")
add_bullet(doc, "Collaboration multi-utilisateurs sur un même manuel (édition temps réel).")
add_bullet(doc, "Intégration CI/CD pour une mise à jour automatique de la documentation lors des déploiements.")
add_bullet(doc, "Application mobile (React Native) pour la consultation des manuels sur smartphone.")
add_bullet(doc, "Tableau de bord analytique pour les administrateurs (statistiques d'utilisation, scores moyens).")

add_para(doc, "DocuSense constitue une première réponse prometteuse à la problématique de la documentation logicielle automatisée et gamifiée. Les technologies utilisées et l'architecture mise en place offrent une base solide et extensible pour des développements futurs.")

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# RÉFÉRENCES BIBLIOGRAPHIQUES (format APA)
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Références bibliographiques", level=1)

references = [
    "Géron, A. (2022). Hands-On Machine Learning with Scikit-Learn, Keras, and TensorFlow (3e éd.). O'Reilly Media.",
    "Google DeepMind. (2024). Gemini: A Family of Highly Capable Multimodal Models. https://deepmind.google/technologies/gemini. Consulté le 10 mars 2025.",
    "Meta AI. (2023). React Documentation — React v19. https://react.dev. Consulté le 15 janvier 2025.",
    "OpenJS Foundation. (2024). Node.js Documentation v20. https://nodejs.org/docs. Consulté le 20 janvier 2025.",
    "Russell, S., & Norvig, P. (2022). Artificial Intelligence: A Modern Approach (4e éd.). Pearson.",
    "Sequelize. (2024). Sequelize ORM Documentation v6. https://sequelize.org/docs/v6. Consulté le 5 février 2025.",
    "Stack Overflow. (2023). Developer Survey 2023. https://survey.stackoverflow.co/2023. Consulté le 1er mars 2025.",
    "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, Ł., & Polosukhin, I. (2017). Attention is All You Need. Advances in Neural Information Processing Systems, 30.",
    "Walbaum, B. (2025). EDUCATION.IA : Réinventer l'éducation à l'ère de l'intelligence artificielle. Débats Publics Éditions.",
]
for ref in references:
    p = doc.add_paragraph()
    para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=3, space_after=3, first_indent=-0.5)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(ref)
    set_font(run, size=11)

page_break(doc)

# ═══════════════════════════════════════════════════════════════
# ANNEXES
# ═══════════════════════════════════════════════════════════════
add_heading(doc, "Annexe A – Instructions de déploiement", level=1)

add_heading(doc, "Prérequis", level=2)
add_bullet(doc, "Node.js v18+ installé")
add_bullet(doc, "MySQL 8.x installé et en cours d'exécution")
add_bullet(doc, "Compte Google AI Studio pour la clé API Gemini")

add_heading(doc, "Installation du backend", level=2)
add_para(doc, "cd backend\nnpm install\n# Configurer les variables dans .env\nnode server.js",
         font_size=10, first_indent=0)

add_heading(doc, "Variables d'environnement (.env)", level=2)
add_para(doc, "DB_HOST=localhost\nDB_NAME=docusense\nDB_USER=root\nDB_PASSWORD=votre_mot_de_passe\nJWT_SECRET=votre_cle_secrete\nGEMINI_API_KEY=votre_cle_api_gemini\nPORT=5000",
         font_size=10, first_indent=0)

add_heading(doc, "Installation du frontend", level=2)
add_para(doc, "cd frontend2\nnpm install\nnpm start   # Lance sur http://localhost:3000",
         font_size=10, first_indent=0)

add_heading(doc, "Création de la base de données", level=2)
add_para(doc, "CREATE DATABASE docusense CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;\n-- Sequelize crée automatiquement les tables au démarrage via sequelize.sync({ alter: true })",
         font_size=10, first_indent=0)

page_break(doc)

add_heading(doc, "Annexe B – Exemple de réponse JSON de l'API Gemini", level=1)
add_para(doc,
'{\n  "steps": [\n    { "title": "Connexion", "description": "Accédez à la page de connexion.", "type": "text" },\n    { "title": "Création du projet", "description": "Cliquez sur + Nouveau manuel.", "type": "text" },\n    { "title": "Paramétrage", "description": "Renseignez le titre et la description.", "type": "text" },\n    { "title": "Génération IA", "description": "Utilisez le champ description pour générer les étapes.", "type": "text" },\n    { "title": "Publication", "description": "Publiez votre manuel pour le rendre accessible.", "type": "text" }\n  ]\n}',
         font_size=10, first_indent=0)

# ═══════════════════════════════════════════════════════════════
# NUMÉROTATION DES PAGES (bas à droite, toutes sections)
# ═══════════════════════════════════════════════════════════════
for i, sec in enumerate(doc.sections):
    configure_section(sec)
    add_page_number_footer(sec)

# ─── Sauvegarde ───────────────────────────────────────────────────────────────
output_path = r"c:\Users\Brothres\Desktop\pfe\docusense\rapport_docusense.docx"
doc.save(output_path)
print(f"Fichier créé : {output_path}")
